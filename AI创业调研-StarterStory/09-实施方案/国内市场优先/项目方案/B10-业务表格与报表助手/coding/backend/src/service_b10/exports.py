from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
import ctypes
import errno
import sys
from pathlib import Path, PurePath, PureWindowsPath
from typing import Any

from openpyxl import Workbook


SCHEMA_VERSION = 1
WORKBOOK_NAME = "report.xlsx"
SNAPSHOT_NAME = "snapshot.json"
BUNDLE_MARKER_NAME = ".service-b10-export.json"
BUNDLE_MARKER_BYTES = b'{"format":"service_b10_export_bundle","schema_version":1}'


def export_bundle(directory: Path, payload: dict) -> Path:
    """Publish one verified, immutable reviewed-version export directory."""
    destination = Path(directory)
    _validate_payload(payload)
    if os.path.lexists(destination):
        raise ValueError("VERSION_EXISTS")

    parent = destination.parent
    parent.mkdir(parents=True, exist_ok=True)
    _reject_symlink_components(parent)
    temporary = Path(tempfile.mkdtemp(prefix=f".{destination.name}-", dir=parent))
    try:
        snapshot_bytes = _canonical_bytes(payload)
        (temporary / SNAPSHOT_NAME).write_bytes(snapshot_bytes)
        (temporary / BUNDLE_MARKER_NAME).write_bytes(BUNDLE_MARKER_BYTES)
        _build_workbook(payload).save(temporary / WORKBOOK_NAME)
        files = [
            _file_record(temporary / SNAPSHOT_NAME),
            _file_record(temporary / WORKBOOK_NAME),
        ]
        manifest_data = {
            "schema_version": SCHEMA_VERSION,
            "client": payload["client"],
            "task": payload["task"],
            "revision": payload["revision"],
            "payload_sha256": _digest(snapshot_bytes),
            "files": files,
        }
        manifest = temporary / "manifest.json"
        manifest.write_bytes(_canonical_bytes(manifest_data))
        verify_bundle(manifest)
        _publish_no_replace(temporary, destination)
        return destination / "manifest.json"
    except BaseException:
        if temporary.exists():
            shutil.rmtree(temporary)
        raise


def verify_bundle(manifest: Path) -> None:
    """Verify every declared file and bind the manifest to its frozen payload."""
    manifest = Path(manifest)
    if manifest.is_symlink() or not manifest.is_file():
        raise ValueError("INVALID_PATH")
    try:
        meta = json.loads(manifest.read_text(encoding="utf-8"))
        files = meta["files"]
    except (OSError, UnicodeError, json.JSONDecodeError, KeyError, TypeError):
        raise ValueError("INVALID_MANIFEST") from None
    if not isinstance(files, list) or not files:
        raise ValueError("INVALID_MANIFEST")

    base = manifest.parent.resolve()
    names: set[str] = set()
    for item in files:
        if not isinstance(item, dict):
            raise ValueError("INVALID_MANIFEST")
        name = item.get("name")
        expected = item.get("sha256")
        if not _safe_file_name(name) or name in names:
            raise ValueError("INVALID_PATH")
        names.add(name)
        path = manifest.parent / name
        if path.is_symlink():
            raise ValueError("INVALID_PATH")
        try:
            resolved = path.resolve(strict=True)
        except (OSError, RuntimeError):
            raise ValueError("DIGEST_MISMATCH") from None
        if not resolved.is_relative_to(base):
            raise ValueError("INVALID_PATH")
        if (
            not path.is_file()
            or not isinstance(expected, str)
            or len(expected) != 64
            or _digest(path.read_bytes()) != expected
        ):
            raise ValueError("DIGEST_MISMATCH")

    if names != {SNAPSHOT_NAME, WORKBOOK_NAME}:
        raise ValueError("INVALID_MANIFEST")
    marker = manifest.parent / BUNDLE_MARKER_NAME
    if marker.exists() and (
        marker.is_symlink()
        or not marker.is_file()
        or marker.read_bytes() != BUNDLE_MARKER_BYTES
    ):
        raise ValueError("INVALID_MANIFEST")
    try:
        snapshot_bytes = (manifest.parent / SNAPSHOT_NAME).read_bytes()
        snapshot = json.loads(snapshot_bytes)
        binding = (
            meta["client"], meta["task"], meta["revision"],
            meta["payload_sha256"],
        )
        expected_binding = (
            snapshot["client"], snapshot["task"], snapshot["revision"],
            _digest(_canonical_bytes(snapshot)),
        )
    except (KeyError, TypeError, UnicodeError, json.JSONDecodeError):
        raise ValueError("BINDING_MISMATCH") from None
    if binding != expected_binding:
        raise ValueError("BINDING_MISMATCH")


def _validate_payload(payload: dict) -> None:
    if not isinstance(payload, dict):
        raise ValueError("INVALID_PAYLOAD")
    state = payload.get("state")
    if state not in {"reviewed", "delivered", "accepted", "returned"}:
        raise ValueError("SNAPSHOT_NOT_REVIEWED")
    if state != "reviewed":
        events = payload.get("events")
        if not isinstance(events, list):
            raise ValueError("SNAPSHOT_NOT_REVIEWED")
        required = [("review", "submitted", "reviewed"),
                    ("deliver", "reviewed", "delivered")]
        if state == "accepted":
            required.append(("confirm", "delivered", "accepted"))
        if state == "returned":
            required.append(("return", "delivered", "returned"))
        position = 0
        for action, before, after in required:
            for index in range(position, len(events)):
                event = events[index]
                actor = event.get("actor") if isinstance(event, dict) else None
                if (isinstance(actor, dict)
                        and all(isinstance(actor.get(k), str) and actor[k].strip()
                                for k in ("name", "channel", "evidence"))
                        and (event.get("action"), event.get("from_state"), event.get("to_state"))
                        == (action, before, after)):
                    position = index + 1
                    break
            else:
                raise ValueError("SNAPSHOT_NOT_REVIEWED")
    if not isinstance(payload.get("client"), str) or not payload["client"]:
        raise ValueError("INVALID_PAYLOAD")
    if not isinstance(payload.get("task"), str) or not payload["task"]:
        raise ValueError("INVALID_PAYLOAD")
    revision = payload.get("revision")
    if isinstance(revision, bool) or not isinstance(revision, int) or revision < 1:
        raise ValueError("INVALID_PAYLOAD")
    report = payload.get("report_snapshot")
    if not isinstance(report, dict) or not isinstance(report.get("metrics"), dict):
        raise ValueError("INVALID_PAYLOAD")
    if not isinstance(payload.get("input_snapshot"), dict):
        raise ValueError("INVALID_PAYLOAD")
    if not isinstance(payload.get("metric_rules"), dict):
        raise ValueError("INVALID_PAYLOAD")


def _build_workbook(payload: dict) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "summary"
    summary.append(["指标", "结果", "币种", "期间", "规则版本"])
    metrics = payload["report_snapshot"]["metrics"]
    for key in ("paid_less_refund", "ads"):
        metric = metrics.get(key)
        if not isinstance(metric, dict):
            raise ValueError("INVALID_PAYLOAD")
        value = metric.get("value")
        display = "不可计算" if value is None else _minor_units(value, metric.get("places"))
        summary.append([
            _safe_text(key), display, _safe_text(metric.get("currency")),
            _safe_text(metric.get("period")), _safe_text(metric.get("rule_version")),
        ])

    status = {"reviewed": "已审阅（待交付）", "delivered": "已交付（待客户确认）",
              "accepted": "客户已确认", "returned": "已退回（历史记录，不是有效终稿）"}
    summary.append(["版本状态", status[payload["state"]], "版本", payload["revision"]])

    sources = workbook.create_sheet("sources")
    sources.append([
        "指标", "类别", "来源文件", "工作表", "行号", "来源文件ID", "批次ID",
        "匹配状态", "异常ID", "排除原因",
    ])
    resolutions = {
        item.get("refund_id"): item
        for item in payload.get("resolutions", []) if isinstance(item, dict)
    }
    for metric_name in sorted(metrics):
        metric = metrics[metric_name]
        if not isinstance(metric, dict):
            raise ValueError("INVALID_PAYLOAD")
        for source in metric.get("sources", []):
            if not isinstance(source, dict):
                raise ValueError("INVALID_PAYLOAD")
            _append_source(sources, metric_name, source, "matched", "", "")

    for detail in payload["report_snapshot"].get("exception_details", []):
        if not isinstance(detail, dict):
            raise ValueError("INVALID_PAYLOAD")
        exception_id = detail.get("id")
        resolution = resolutions.get(exception_id, {})
        exclusion_reason = (
            resolution.get("explanation")
            or resolution.get("resolution")
            or detail.get("code")
        )
        for source in detail.get("sources", []):
            if not isinstance(source, dict):
                raise ValueError("INVALID_PAYLOAD")
            _append_source(
                sources, "exception", source, "unmatched", exception_id,
                exclusion_reason,
            )

    exceptions = workbook.create_sheet("exceptions")
    exceptions.append([
        "异常代码", "退款ID", "店铺", "订单", "金额", "处理结果", "说明"
    ])
    for detail in payload["report_snapshot"].get("exception_details", []):
        if not isinstance(detail, dict):
            raise ValueError("INVALID_PAYLOAD")
        resolution = resolutions.get(detail.get("id"), {})
        places = payload["report_snapshot"].get("places", 2)
        exceptions.append([
            _safe_text(detail.get("code")), _safe_text(detail.get("id")),
            _safe_text(detail.get("shop")), _safe_text(detail.get("order")),
            _minor_units(detail.get("amount"), places),
            _safe_text(resolution.get("resolution")),
            _safe_text(resolution.get("explanation")),
        ])

    rules = workbook.create_sheet("rules")
    rules.append(["字段", "值"])
    for key in sorted(payload["metric_rules"]):
        value = payload["metric_rules"][key]
        rendered = value if isinstance(value, (str, int, float)) else json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        )
        rules.append([_safe_text(key), _safe_text(rendered)])
    return workbook


def _minor_units(value: Any, places: Any) -> str:
    if isinstance(value, bool) or not isinstance(value, int):
        raise ValueError("INVALID_PAYLOAD")
    if isinstance(places, bool) or not isinstance(places, int) or not 0 <= places <= 9:
        raise ValueError("INVALID_PAYLOAD")
    sign = "-" if value < 0 else ""
    digits = str(abs(value))
    if places == 0:
        return sign + digits
    digits = digits.zfill(places + 1)
    return f"{sign}{digits[:-places]}.{digits[-places:]}"


def _append_source(sheet, label, source, match_status, exception_id, reason) -> None:
    row_index = source.get("row_index")
    if not isinstance(row_index, int) or isinstance(row_index, bool):
        row_index = _safe_text(row_index)
    sheet.append([
        _safe_text(label), _safe_text(source.get("category")),
        _safe_text(_source_basename(source.get("source_file"))),
        _safe_text(source.get("sheet")), row_index,
        _safe_text(source.get("source_file_id")),
        _safe_text(source.get("batch_id")), _safe_text(match_status),
        _safe_text(exception_id), _safe_text(reason),
    ])


def _safe_text(value: Any) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _safe_file_name(value: Any) -> bool:
    if not isinstance(value, str) or not value or "\\" in value:
        return False
    path = PurePath(value)
    return not path.is_absolute() and len(path.parts) == 1 and path.name == value


def _source_basename(value: Any) -> str:
    text = "" if value is None else str(value)
    if "\\" in text:
        return PureWindowsPath(text).name
    return Path(text).name


def _reject_symlink_components(path: Path) -> None:
    if path.is_symlink():
        raise ValueError("INVALID_PATH")


def _publish_no_replace(source: Path, destination: Path) -> None:
    """Atomically rename a directory only when destination does not exist."""
    library = ctypes.CDLL(None, use_errno=True)
    source_bytes = os.fsencode(source)
    destination_bytes = os.fsencode(destination)
    if sys.platform == "darwin" and hasattr(library, "renamex_np"):
        operation = library.renamex_np
        operation.argtypes = [ctypes.c_char_p, ctypes.c_char_p, ctypes.c_uint]
        operation.restype = ctypes.c_int
        result = operation(source_bytes, destination_bytes, 0x00000004)
    elif sys.platform.startswith("linux") and hasattr(library, "renameat2"):
        operation = library.renameat2
        operation.argtypes = [
            ctypes.c_int, ctypes.c_char_p, ctypes.c_int, ctypes.c_char_p,
            ctypes.c_uint,
        ]
        operation.restype = ctypes.c_int
        result = operation(-100, source_bytes, -100, destination_bytes, 1)
    else:
        raise RuntimeError("ATOMIC_NOREPLACE_UNSUPPORTED")
    if result == 0:
        return
    error = ctypes.get_errno()
    if error in {errno.EEXIST, errno.ENOTEMPTY}:
        raise ValueError("VERSION_EXISTS")
    raise OSError(error, os.strerror(error), os.fspath(destination))


def _file_record(path: Path) -> dict[str, str]:
    return {"name": path.name, "sha256": _digest(path.read_bytes())}


def _canonical_bytes(value: Any) -> bytes:
    try:
        return json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
            allow_nan=False,
        ).encode("utf-8")
    except (TypeError, ValueError):
        raise ValueError("INVALID_PAYLOAD") from None


def _digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()
