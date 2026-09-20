import hashlib
import json
import tempfile
from copy import deepcopy
from decimal import localcontext
from pathlib import Path

import pytest
from openpyxl import load_workbook

from service_b10.exports import export_bundle, verify_bundle


FIXTURE = Path("tests/fixtures/approved_payload.json")


def _payload():
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def _workbook(manifest):
    info = json.loads(manifest.read_text(encoding="utf-8"))
    workbook = next(item for item in info["files"] if item["name"].endswith(".xlsx"))
    return load_workbook(manifest.parent / workbook["name"], data_only=False)


def test_export_manifest_tampering(tmp_path):
    manifest = export_bundle(tmp_path / "v1", _payload())
    assert json.loads((manifest.parent / ".service-b10-export.json").read_text()) == {
        "format": "service_b10_export_bundle", "schema_version": 1
    }
    verify_bundle(manifest)
    info = json.loads(manifest.read_text())
    assert (info["client"], info["task"], info["revision"]) == ("a", "t1", 1)
    file = manifest.parent / info["files"][0]["name"]
    file.write_bytes(b"tampered")
    with pytest.raises(ValueError, match="^DIGEST_MISMATCH$"):
        verify_bundle(manifest)


def test_workbook_has_literal_minor_unit_amounts_and_provenance(tmp_path):
    manifest = export_bundle(tmp_path / "v1", _payload())
    workbook = _workbook(manifest)
    assert workbook.sheetnames == ["summary", "sources", "exceptions", "rules"]
    summary = workbook["summary"]
    assert summary["B2"].value == "900.00"
    assert summary["B3"].value == "50.00"
    assert summary["B2"].data_type != "f"
    sources = list(workbook["sources"].values)
    assert any(
        row[:7] == ("paid_less_refund", "orders", "orders.xlsx", "Orders", 2,
                    "orders-source-1", "orders-batch-1")
        for row in sources
    )
    assert all("/private/" not in str(cell) for row in sources for cell in row)


def test_windows_source_paths_are_reduced_to_human_basename(tmp_path):
    payload = _payload()
    payload["report_snapshot"]["metrics"]["ads"]["sources"][0][
        "source_file"
    ] = r"C:\customer\private\ads.xlsx"
    workbook = _workbook(export_bundle(tmp_path / "v1", payload))
    source_files = [row[2] for row in list(workbook["sources"].values)[1:]]
    assert "ads.xlsx" in source_files
    assert all("\\" not in name and "/" not in name for name in source_files)


def test_missing_ads_is_explicitly_uncomputable(tmp_path):
    payload = _payload()
    payload["input_snapshot"]["ads"] = None
    payload["report_snapshot"]["ads"] = None
    payload["report_snapshot"]["metrics"]["ads"].update(
        value=None, reason="MISSING_ADS"
    )
    manifest = export_bundle(tmp_path / "v1", payload)
    assert _workbook(manifest)["summary"]["B3"].value == "不可计算"


@pytest.mark.parametrize("minor_units,expected", [
    (1234567890123456789012345678901, "12345678901234567890123456789.01"),
    (-12345, "-123.45"),
])
def test_money_format_is_exact_for_all_integers_and_ambient_precision(
    tmp_path, minor_units, expected
):
    payload = _payload()
    payload["report_snapshot"]["paid_less_refund"] = minor_units
    payload["report_snapshot"]["metrics"]["paid_less_refund"]["value"] = minor_units
    with localcontext() as context:
        context.prec = 6
        manifest = export_bundle(tmp_path / "v1", payload)
    assert _workbook(manifest)["summary"]["B2"].value == expected
    verify_bundle(manifest)


def test_manifest_binds_identity_and_canonical_payload_digest(tmp_path):
    payload = _payload()
    manifest = export_bundle(tmp_path / "unrelated-name", payload)
    info = json.loads(manifest.read_text(encoding="utf-8"))
    expected = hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True,
                   separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    assert info["payload_sha256"] == expected
    snapshot = json.loads((manifest.parent / "snapshot.json").read_text())
    assert snapshot == payload
    verify_bundle(manifest)


@pytest.mark.parametrize("field,value", [
    ("client", "other"), ("task", "other"), ("revision", 2),
    ("payload_sha256", "0" * 64),
])
def test_verify_rejects_manifest_binding_tampering(tmp_path, field, value):
    manifest = export_bundle(tmp_path / "v1", _payload())
    info = json.loads(manifest.read_text())
    info[field] = value
    manifest.write_text(json.dumps(info), encoding="utf-8")
    with pytest.raises(ValueError, match="^BINDING_MISMATCH$"):
        verify_bundle(manifest)


@pytest.mark.parametrize("name", ["../outside", "/tmp/outside", "snapshot.json/../x"])
def test_verify_rejects_unsafe_file_paths(tmp_path, name):
    manifest = export_bundle(tmp_path / "v1", _payload())
    info = json.loads(manifest.read_text())
    info["files"][0]["name"] = name
    manifest.write_text(json.dumps(info), encoding="utf-8")
    with pytest.raises(ValueError, match="^INVALID_PATH$"):
        verify_bundle(manifest)


def test_verify_rejects_symlink_file_escape(tmp_path):
    manifest = export_bundle(tmp_path / "v1", _payload())
    outside = tmp_path / "outside"
    outside.write_bytes(b"outside")
    info = json.loads(manifest.read_text())
    item = info["files"][0]
    path = manifest.parent / item["name"]
    path.unlink()
    path.symlink_to(outside)
    item["sha256"] = hashlib.sha256(b"outside").hexdigest()
    manifest.write_text(json.dumps(info), encoding="utf-8")
    with pytest.raises(ValueError, match="^INVALID_PATH$"):
        verify_bundle(manifest)


def test_existing_version_is_never_overwritten(tmp_path):
    destination = tmp_path / "v1"
    manifest = export_bundle(destination, _payload())
    before = manifest.read_bytes()
    with pytest.raises(ValueError, match="^VERSION_EXISTS$"):
        export_bundle(destination, _payload())
    assert manifest.read_bytes() == before


@pytest.mark.parametrize("occupied", [False, True])
def test_concurrent_destination_is_atomically_preserved(
    tmp_path, monkeypatch, occupied
):
    import service_b10.exports as exports

    destination = tmp_path / "v1"
    original_publish = exports._publish_no_replace

    def preempt(temporary, target):
        target.mkdir()
        if occupied:
            (target / "owner.txt").write_text("other publisher", encoding="utf-8")
        inode = target.stat().st_ino
        try:
            return original_publish(temporary, target)
        finally:
            assert target.stat().st_ino == inode

    monkeypatch.setattr(exports, "_publish_no_replace", preempt)
    with pytest.raises(ValueError, match="^VERSION_EXISTS$"):
        export_bundle(destination, _payload())
    assert destination.is_dir()
    if occupied:
        assert (destination / "owner.txt").read_text() == "other publisher"
    assert not any(path.name.startswith(".v1-") for path in tmp_path.iterdir())


def test_unreviewed_payload_is_rejected_without_publication(tmp_path):
    payload = _payload()
    payload["state"] = "draft"
    destination = tmp_path / "v1"
    with pytest.raises(ValueError, match="^SNAPSHOT_NOT_REVIEWED$"):
        export_bundle(destination, payload)
    assert not destination.exists()


def test_generation_failure_leaves_no_version_or_temp_directory(tmp_path, monkeypatch):
    import service_b10.exports as exports

    def fail_save(self, filename):
        raise OSError("disk full")

    monkeypatch.setattr(exports.Workbook, "save", fail_save)
    destination = tmp_path / "v1"
    with pytest.raises(OSError, match="disk full"):
        export_bundle(destination, _payload())
    assert not destination.exists()
    assert list(tmp_path.iterdir()) == []


def test_standard_temporary_directory_path_is_supported():
    with tempfile.TemporaryDirectory() as temporary:
        manifest = export_bundle(Path(temporary) / "v1", _payload())
        verify_bundle(manifest)


def test_untrusted_spreadsheet_text_is_formula_sanitized(tmp_path):
    payload = _payload()
    payload["report_snapshot"]["metrics"]["ads"]["sources"][0].update(
        source_file="=WEBSERVICE(\"https://bad\")",
        sheet="+cmd|' /C calc'!A0",
    )
    payload["metric_rules"]["note"] = "@SUM(1,1)"
    payload["report_snapshot"]["exception_details"] = [{
        "code": "UNMATCHED_REFUND", "id": "-1+1", "shop": "a",
        "order": "1", "amount": 1, "sources": [],
    }]
    workbook = _workbook(export_bundle(tmp_path / "v1", payload))
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f"
                if isinstance(cell.value, str):
                    assert not cell.value.startswith(("=", "+", "-", "@"))


def test_unmatched_refund_sources_are_exported_with_exclusion_lineage(tmp_path):
    payload = _payload()
    payload["report_snapshot"]["exception_details"] = [{
        "code": "UNMATCHED_REFUND",
        "id": "rx",
        "shop": "a",
        "order": "missing",
        "amount": 100,
        "sources": [
            {
                "category": "refunds",
                "source_file": "/customer/unmatched-a.xlsx",
                "sheet": "Refunds-A",
                "row_index": 42,
                "source_file_id": "exception-source-a",
                "batch_id": "exception-batch-a",
            },
            {
                "category": "refunds",
                "source_file": r"C:\customer\unmatched-b.xlsx",
                "sheet": "@Refunds-B",
                "row_index": 43,
                "source_file_id": "exception-source-b",
                "batch_id": "exception-batch-b",
            },
        ],
    }]
    payload["report_snapshot"]["exceptions"] = ["rx"]
    payload["resolutions"] = [{
        "refund_id": "rx",
        "resolution": "excluded",
        "explanation": "Order not supplied",
        "actor": {},
        "evidence": "fictional",
    }]
    workbook = _workbook(export_bundle(tmp_path / "v1", payload))
    rows = list(workbook["sources"].values)
    assert rows[0][-3:] == ("匹配状态", "异常ID", "排除原因")
    exception_rows = [row for row in rows[1:] if row[-2] == "rx"]
    assert len(exception_rows) == 2
    assert {row[2] for row in exception_rows} == {
        "unmatched-a.xlsx", "unmatched-b.xlsx"
    }
    assert {row[5] for row in exception_rows} == {
        "exception-source-a", "exception-source-b"
    }
    assert all(row[-3:] == ("unmatched", "rx", "Order not supplied")
               for row in exception_rows)
    assert exception_rows[1][3].startswith("'")
