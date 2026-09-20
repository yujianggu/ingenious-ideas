"""CLI acceptance: real subprocesses, synthetic inputs, no provider accounts."""
import json
import os
from pathlib import Path
import subprocess
import sys

import pytest

from service_b10.store import Store
from service_b10.inputs import get_active_import
from service_b10.workflow import get_revision
from service_b10.exports import verify_bundle, export_bundle


def cli(*args, code=0):
    result = subprocess.run([sys.executable, '-m', 'service_b10.cli', *map(str, args)],
                            capture_output=True, text=True, timeout=15)
    assert result.returncode == code, (result.stdout, result.stderr)
    return json.loads(result.stdout) if code == 0 else result.stderr


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
    return path


def setup_job(tmp_path, client='alice'):
    root = tmp_path / 'live'
    scope = ['--root', root, '--client', client, '--task', 'week1']
    cli('new', *scope)
    source = tmp_path / 'evidence.txt'
    source.write_text('虚构人工核对原始邮件，已逐项复核。')
    digest = cli('evidence', *scope, '--input', source)['evidence']
    actor = write_json(tmp_path / f'{client}-actor.json',
                       {'name': '虚构交付员', 'channel': 'operator', 'evidence': digest})
    for category, contents in {
        'orders': '店铺,订单号,实收\na,1,1000.00\n',
        'refunds': '退款号,店铺,订单号,退款金额\nr1,a,1,100.00\nr2,a,9,2.00\n',
        'ads': '广告费\n50.00\n',
    }.items():
        path = tmp_path / f'{client}-{category}.csv'
        path.write_text(contents)
        cli('submit', *scope, '--category', category, '--period', '2026-09',
            '--input', path, '--schema', 'config/schema.example.json')
    return root, scope, actor, digest


def ready(scope, actor, tmp_path, revision):
    cli('produce', *scope, '--period', '2026-09', '--actor-record', actor)
    cli('review', *scope, '--revision', revision, '--actor-record', actor)
    cli('deliver', *scope, '--revision', revision, '--actor-record', actor, code=2)
    rules = write_json(tmp_path / 'rules.json', {'paid_less_refund': '本期实收减匹配退款；不称净利润',
                                              'ads': '缺失不可计算，非零填补'})
    cli('freeze-rules', *scope, '--revision', revision, '--actor-record', actor, '--rules', rules)
    cli('deliver', *scope, '--revision', revision, '--actor-record', actor, code=2)
    cli('resolve', *scope, '--revision', revision, '--actor-record', actor,
        '--refund-id', 'r2', '--resolution', 'exclude', '--reason', '跨期退款，保留异常不计入')
    return cli('deliver', *scope, '--revision', revision, '--actor-record', actor)


def receipt(scope, delivery, tmp_path, evidence, revision):
    value = {'evidence_kind': 'customer_confirmation_receipt', 'purpose': 'confirm_delivery',
             'client': 'alice', 'task': 'week1', 'revision': revision,
             'delivery_instance_id': delivery['delivery_instance_id'],
             'snapshot_digest': delivery['snapshot_digest'],
             'source_channel': 'email', 'source_evidence': evidence}
    path = write_json(tmp_path / f'receipt-{revision}.json', value)
    digest = cli('receipt', *scope, '--revision', revision, '--input', path)['evidence']
    return write_json(tmp_path / f'customer-{revision}.json',
                      {'name': '虚构客户复核人', 'channel': 'customer', 'evidence': digest})


def test_cli_complete_history_two_clients_external_export_backup_restore(tmp_path):
    cli('init', '--root', tmp_path / 'live')
    root, scope, actor, evidence = setup_job(tmp_path)
    _, other, other_actor, _ = setup_job(tmp_path, 'bob')
    cli('produce', *other, '--period', '2026-09', '--actor-record', other_actor)
    first = ready(scope, actor, tmp_path, 1)
    old_customer = receipt(scope, first, tmp_path, evidence, 1)
    cli('return', *scope, '--revision', 1, '--actor-record', old_customer, '--reason', '实收漏单')
    cli('confirm', *scope, '--revision', 1, '--actor-record', old_customer, code=2)
    cli('amend', *scope, '--revision', 1, '--actor-record', actor, '--reason', '补全原件')
    batch = get_active_import(Store(root), 'alice', 'week1', 'orders', '2026-09')
    replacement = tmp_path / '修订订单.csv'
    replacement.write_text('店铺,订单号,实收\na,1,2000.00\n')
    cli('submit', *scope, '--category', 'orders', '--period', '2026-09', '--input', replacement,
        '--schema', 'config/schema.example.json', '--supersedes', batch.batch_id)
    second = ready(scope, actor, tmp_path, 2)
    assert second['delivery_instance_id'] != first['delivery_instance_id']
    assert second['snapshot_digest'] != first['snapshot_digest']
    cli('confirm', *scope, '--revision', 2, '--actor-record', old_customer, code=2)
    plain_customer = write_json(tmp_path / 'plain.json', {'name': '任意姓名', 'channel': 'customer', 'evidence': evidence})
    cli('confirm', *scope, '--revision', 2, '--actor-record', plain_customer, code=2)
    customer = receipt(scope, second, tmp_path, evidence, 2)
    cli('confirm', *scope, '--revision', 2, '--actor-record', customer)
    output = tmp_path / 'delivery-v2'
    cli('export', *scope, '--revision', 2, '--out', output)
    cli('export', *scope, '--revision', 1, '--out', tmp_path / 'delivery-v1')
    verify_bundle(output / 'manifest.json')
    assert json.loads((output / 'snapshot.json').read_text())['state'] == 'accepted'
    cli('backup', '--root', root, '--out', tmp_path / 'backup')
    meta = json.loads((tmp_path / 'backup/manifest.json').read_text())
    assert len(meta['bundles']) == 2
    cli('restore', '--manifest', tmp_path / 'backup/manifest.json', '--out', tmp_path / 'restored')
    restored = Store(tmp_path / 'restored')
    for directory in meta['bundles']:
        verify_bundle(restored.root / directory / 'manifest.json')
    old = get_revision(restored, 'alice', 'week1', 1)
    latest = get_revision(restored, 'alice', 'week1', 2)
    assert old['report_snapshot']['paid_less_refund'] == 90000
    assert old['state'] == 'returned'
    assert len(old['confirmation_proofs']) == 1
    assert latest['report_snapshot']['paid_less_refund'] == 190000
    assert latest['state'] == 'accepted'
    assert [e['action'] for e in latest['events']] == ['submit', 'review', 'deliver', 'confirm']
    assert get_revision(restored, 'bob', 'week1')['report_snapshot']['paid_less_refund'] == 90000
    assert get_revision(restored, 'bob', 'week1')['state'] == 'submitted'


def test_cli_invalid_input_preserved_and_logs_are_codes(tmp_path):
    root, scope, actor, _ = setup_job(tmp_path)
    path = tmp_path / 'private-客户-sk-secret.csv'
    data = '店铺,订单号,实收\na,1,sk-secret\n'
    path.write_text(data)
    error = cli('submit', *scope, '--category', 'orders', '--period', '2026-09',
                '--schema', 'config/schema.example.json', '--input', path, code=2)
    assert 'sk-secret' not in error and str(tmp_path) not in error and '客户' not in error
    assert 'INVALID_AMOUNT' in error
    import hashlib
    assert Store(root).read('alice', 'week1', hashlib.sha256(data.encode()).hexdigest()) == data.encode()
    assert get_active_import(Store(root), 'alice', 'week1', 'orders', '2026-09').rows[0]['paid'] == 100000
    cli('evidence', *scope, '--input', tmp_path / 'missing-secret', code=1)
    error = cli('new', '--root', root, '--client', '../secret', '--task', 'secret/path', code=2)
    assert 'secret' not in error


def test_help_parser_redacts_raw_arguments():
    result = subprocess.run([sys.executable, '-m', 'service_b10.cli', '--help'], capture_output=True, text=True)
    assert result.returncode == 0 and 'backup' in result.stdout
    error = cli('submit', '--sk-secret', code=2)
    assert 'sk-secret' not in error


def test_export_accepted_without_review_event_is_rejected(tmp_path):
    payload = json.loads(Path('tests/fixtures/approved_payload.json').read_text())
    payload['state'] = 'accepted'
    payload['events'] = []
    with pytest.raises(ValueError, match='SNAPSHOT_NOT_REVIEWED'):
        export_bundle(tmp_path / 'out', payload)


def model_payload():
    return {'config': {'endpoint': 'https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions',
                       'model_id': 'explicit-evaluation-id', 'evaluation_record': 'local-evaluation-1',
                       'timeout_seconds': 1, 'max_tokens': 128, 'max_total_tokens': 2000, 'retries': 2},
            'input_snapshot': {'orders': [{'shop': 's', 'order': '1', 'paid': 100000}], 'refunds': [], 'ads': None},
            'source_ids': ['source1'],
            'metrics': {'paid_less_refund': 100000, 'ads': None}}


def candidate_response(**changes):
    return {'text': '实收减退款待人工核对', 'source_ids': ['source1'],
            'metrics': {'paid_less_refund': 100000, 'ads': None},
            'usage': {'total_tokens': 100}, **changes}


def test_model_disabled_and_unauthorized_never_transport(monkeypatch):
    from service_b10.model import generate_candidate
    def forbidden(_):
        raise AssertionError('network must not run')
    assert generate_candidate({}, enabled=False, permitted=False, transport=forbidden) is None
    with pytest.raises(ValueError, match='MODEL_NOT_AUTHORIZED'):
        generate_candidate({}, enabled=True, permitted=False, transport=forbidden)


@pytest.mark.parametrize('response,code', [
    ('not json', 'INVALID_MODEL_RESPONSE'),
    (candidate_response(source_ids=['unknown']), 'INVALID_MODEL_SOURCE'),
    (candidate_response(metrics={'paid_less_refund': 999, 'ads': None}), 'INVALID_MODEL_METRICS'),
    (candidate_response(usage={'total_tokens': 99999}), 'MODEL_USAGE_LIMIT'),
    (candidate_response(state='reviewed'), 'INVALID_MODEL_RESPONSE'),
])
def test_model_validation(response, code, monkeypatch):
    from service_b10.model import generate_candidate
    monkeypatch.setenv('DASHSCOPE_API_KEY', 'sk-secret')
    with pytest.raises(ValueError, match=f'^{code}$'):
        generate_candidate(model_payload(), enabled=True, permitted=True, transport=lambda _: response)


def test_model_timeout_bounded_and_no_key(monkeypatch, capsys):
    from service_b10.model import generate_candidate
    monkeypatch.delenv('DASHSCOPE_API_KEY', raising=False)
    with pytest.raises(ValueError, match='MODEL_KEY_REQUIRED'):
        generate_candidate(model_payload(), enabled=True, permitted=True, transport=lambda _: candidate_response())
    monkeypatch.setenv('DASHSCOPE_API_KEY', 'sk-secret')
    attempts = []
    def timeout(request):
        attempts.append(request)
        raise TimeoutError('sk-secret /private/customer')
    with pytest.raises(ValueError, match='^MODEL_TIMEOUT$'):
        generate_candidate(model_payload(), enabled=True, permitted=True, transport=timeout)
    assert len(attempts) == 3
    assert capsys.readouterr() == ('', '')


def test_model_good_candidate_never_reviewed_or_sends_raw_identifiers(monkeypatch):
    from service_b10.model import generate_candidate
    monkeypatch.setenv('DASHSCOPE_API_KEY', 'sk-secret')
    def transport(request):
        assert 'input_snapshot' not in request and 'config' not in request
        assert 'explicit-evaluation-id' == request['model']
        assert 'sk-secret' not in json.dumps(request)
        assert '"shop"' not in json.dumps(request)
        return candidate_response()
    result = generate_candidate(model_payload(), enabled=True, permitted=True, transport=transport)
    assert result['state'] == 'candidate'
    assert result['metrics']['paid_less_refund'] == 100000


@pytest.mark.parametrize('endpoint', ['http://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions',
    'https://dashscope-intl.aliyuncs.com/compatible-mode/v1/chat/completions',
    'https://dashscope.aliyuncs.com.evil.test/compatible-mode/v1/chat/completions',
    'https://secret@dashscope.aliyuncs.com/compatible-mode/v1/chat/completions'])
def test_model_endpoint_allowlist(endpoint, monkeypatch):
    from service_b10.model import generate_candidate
    monkeypatch.setenv('DASHSCOPE_API_KEY', 'sk-secret')
    payload = model_payload()
    payload['config']['endpoint'] = endpoint
    with pytest.raises(ValueError, match='INVALID_MODEL_CONFIG'):
        generate_candidate(payload, enabled=True, permitted=True, transport=lambda _: candidate_response())


@pytest.fixture(autouse=True)
def forbid_network_for_cli_subprocesses(tmp_path, monkeypatch):
    guard = tmp_path / 'network-guard'
    guard.mkdir()
    (guard / 'sitecustomize.py').write_text(
        'import socket\n'
        'def deny(*a, **k):\n    raise AssertionError("NETWORK_FORBIDDEN")\n'
        'socket.create_connection = deny\nsocket.socket.connect = deny\n')
    monkeypatch.setenv('PYTHONPATH', str(guard) + os.pathsep + os.environ.get('PYTHONPATH', ''))


def test_http_redirect_and_nonjson_fail_without_following(monkeypatch):
    import service_b10.model as model
    monkeypatch.setenv('DASHSCOPE_API_KEY', 'sk-secret')
    connections = []
    class Response:
        status = 302
        def read(self, limit):
            return b'not-json sk-secret'
    class Connection:
        def __init__(self, host, timeout):
            connections.append((host, timeout))
        def request(self, method, path, body, headers):
            assert method == 'POST'
            assert headers['Authorization'] == 'Bearer sk-secret'
        def getresponse(self):
            return Response()
        def close(self):
            pass
    monkeypatch.setattr(model.http.client, 'HTTPSConnection', Connection)
    with pytest.raises(ValueError, match='MODEL_REDIRECT_REJECTED'):
        model.generate_candidate(model_payload(), enabled=True, permitted=True)
    assert connections == [('dashscope.aliyuncs.com', 1)]
    Response.status = 200
    with pytest.raises(ValueError, match='INVALID_MODEL_RESPONSE'):
        model.generate_candidate(model_payload(), enabled=True, permitted=True)


def test_model_payload_and_usage_limits_before_transport(monkeypatch):
    from service_b10.model import generate_candidate
    monkeypatch.setenv('DASHSCOPE_API_KEY', 'sk-secret')
    def forbidden(_):
        raise AssertionError('must reject before transport')
    payload = model_payload()
    payload['source_ids'] = [str(i).zfill(100) for i in range(200)]
    with pytest.raises(ValueError, match='MODEL_PAYLOAD_LIMIT'):
        generate_candidate(payload, enabled=True, permitted=True, transport=forbidden)
    payload = model_payload()
    payload['config']['max_total_tokens'] = 1
    with pytest.raises(ValueError, match='MODEL_USAGE_LIMIT'):
        generate_candidate(payload, enabled=True, permitted=True, transport=forbidden)


def test_model_enabled_without_authorization_cli_keeps_inputs(tmp_path):
    root, scope, actor, _ = setup_job(tmp_path)
    error = cli('produce', *scope, '--period', '2026-09', '--actor-record', actor, '--enable-model', code=2)
    assert 'MODEL_NOT_AUTHORIZED' in error
    assert get_active_import(Store(root), 'alice', 'week1', 'orders', '2026-09').rows[0]['paid'] == 100000
    with pytest.raises(ValueError, match='NOT_FOUND'):
        get_revision(Store(root), 'alice', 'week1')


def test_historical_export_displays_returned_status(tmp_path):
    from openpyxl import load_workbook
    payload = json.loads(Path('tests/fixtures/approved_payload.json').read_text())
    actor = payload['events'][0]['actor']
    payload['state'] = 'returned'
    payload['events'] += [
        {'action': 'deliver', 'from_state': 'reviewed', 'to_state': 'delivered', 'actor': actor},
        {'action': 'return', 'from_state': 'delivered', 'to_state': 'returned', 'actor': actor},
    ]
    export_bundle(tmp_path / 'out', payload)
    book = load_workbook(tmp_path / 'out/report.xlsx')
    assert any('已退回' in str(cell.value) for row in book['summary'] for cell in row)


def test_manual_candidate_kept_separate_and_submission_frozen(tmp_path):
    root, scope, actor, _ = setup_job(tmp_path)
    store = Store(root)
    source = get_active_import(store, 'alice', 'week1', 'orders', '2026-09').source_file_id
    candidate = write_json(tmp_path / 'candidate.json', {
        'text': '人工候选解释待审阅', 'source_ids': [source],
        'metrics': {'paid_less_refund': 90000, 'ads': 5000}, 'usage': {'total_tokens': 0}})
    cli('produce', *scope, '--period', '2026-09', '--actor-record', actor, '--candidate', candidate)
    revision = get_revision(store, 'alice', 'week1')
    assert revision['state'] == 'submitted'
    saved = json.loads(store.read('alice', 'week1', revision['input_snapshot']['candidate_evidence']))
    assert saved['state'] == 'candidate'
    assert 'text' not in revision['report_snapshot']
    cli('submit', *scope, '--category', 'orders', '--period', '2026-09',
        '--input', tmp_path / 'alice-orders.csv', '--schema', 'config/schema.example.json', code=2)
    assert get_revision(store, 'alice', 'week1') == revision


def test_cross_client_evidence_does_not_review(tmp_path):
    root, scope, actor, _ = setup_job(tmp_path)
    cli('produce', *scope, '--period', '2026-09', '--actor-record', actor)
    store = Store(root)
    store.create('bob', 'week1')
    digest = store.put('bob', 'week1', 'proof.txt', b'bob-only-evidence')
    other_actor = write_json(tmp_path / 'bob-only.json', {'name': 'reviewer', 'channel': 'operator', 'evidence': digest})
    cli('review', *scope, '--revision', 1, '--actor-record', other_actor, code=2)
    assert get_revision(store, 'alice', 'week1')['state'] == 'submitted'


def test_cli_no_key_logs_no_configuration_or_source(tmp_path, monkeypatch):
    root, scope, actor, evidence = setup_job(tmp_path)
    config = write_json(tmp_path / 'private-config.json', model_payload()['config'])
    auth = write_json(tmp_path / 'private-authorization.json', {
        'purpose': 'authorize_model', 'model_id': 'explicit-evaluation-id',
        'endpoint': model_payload()['config']['endpoint'],
        'allowed_fields': ['metrics', 'source_ids'], 'evidence': evidence})
    monkeypatch.delenv('DASHSCOPE_API_KEY', raising=False)
    error = cli('produce', *scope, '--period', '2026-09', '--actor-record', actor,
                '--enable-model', '--model-config', config, '--model-authorization', auth, code=2)
    assert json.loads(error) == {'code': 'MODEL_KEY_REQUIRED', 'task': 'week1'}


def test_enabled_candidate_retains_evaluation_config(tmp_path, monkeypatch, capsys):
    import service_b10.cli as module
    from service_b10.model import generate_candidate
    root, scope, actor, evidence = setup_job(tmp_path)
    config = write_json(tmp_path / 'config.json', model_payload()['config'])
    auth = write_json(tmp_path / 'auth.json', {
        'purpose': 'authorize_model', 'model_id': 'explicit-evaluation-id',
        'endpoint': model_payload()['config']['endpoint'],
        'allowed_fields': ['metrics', 'source_ids'], 'evidence': evidence})
    monkeypatch.setenv('DASHSCOPE_API_KEY', 'sk-secret')
    def injected(payload, **options):
        options['transport'] = lambda _: candidate_response(metrics=payload['metrics'], source_ids=payload['source_ids'])
        return generate_candidate(payload, **options)
    monkeypatch.setattr(module, 'generate_candidate', injected)
    assert module.main(list(map(str, ['produce', *scope, '--period', '2026-09', '--actor-record', actor,
                          '--enable-model', '--model-config', config, '--model-authorization', auth]))) == 0
    store = Store(root)
    frozen = get_revision(store, 'alice', 'week1')['input_snapshot']
    recorded = json.loads(store.read('alice', 'week1', frozen['model_config_evidence']))
    assert recorded['evaluation_record'] == 'local-evaluation-1'
    assert 'sk-secret' not in capsys.readouterr().out
    assert 'sk-secret' not in json.dumps(recorded)


def test_config_rejects_inline_key_instead_of_persisting_it(monkeypatch):
    from service_b10.model import generate_candidate
    monkeypatch.setenv('DASHSCOPE_API_KEY', 'sk-env-secret')
    payload = model_payload()
    payload['config']['api_key'] = 'sk-inline-secret'
    with pytest.raises(ValueError, match='INVALID_MODEL_CONFIG'):
        generate_candidate(payload, enabled=True, permitted=True, transport=lambda _: candidate_response())


def test_same_revision_redelivery_requires_new_receipt(tmp_path):
    root, scope, actor, evidence = setup_job(tmp_path)
    first = ready(scope, actor, tmp_path, 1)
    customer = receipt(scope, first, tmp_path, evidence, 1)
    cli('return', *scope, '--revision', 1, '--actor-record', customer, '--reason', '重新核对')
    cli('resubmit', *scope, '--revision', 1, '--actor-record', actor)
    cli('review', *scope, '--revision', 1, '--actor-record', actor)
    second = cli('deliver', *scope, '--revision', 1, '--actor-record', actor)
    assert first['delivery_instance_id'] != second['delivery_instance_id']
    assert first['snapshot_digest'] != second['snapshot_digest']
    cli('confirm', *scope, '--revision', 1, '--actor-record', customer, code=2)
    customer = receipt(scope, second, tmp_path, evidence, 1)
    cli('confirm', *scope, '--revision', 1, '--actor-record', customer)
    revision = get_revision(Store(root), 'alice', 'week1')
    assert revision['state'] == 'accepted'
    assert len(revision['confirmation_proofs']) == 2


@pytest.mark.parametrize('kind', ['not_zip', 'invalid_xml'])
def test_corrupt_xlsx_is_domain_error_without_traceback(tmp_path, kind):
    root = tmp_path / 'root'
    Store(root).create('alice', 'week1')
    path = tmp_path / 'secret-customer.xlsx'
    if kind == 'not_zip':
        path.write_bytes(b'not an XLSX zip archive')
    else:
        import io
        import zipfile
        from openpyxl import Workbook
        buffer = io.BytesIO()
        Workbook().save(buffer)
        with zipfile.ZipFile(buffer) as original, zipfile.ZipFile(path, 'w') as broken:
            for name in original.namelist():
                broken.writestr(name, b'broken<xml' if name == 'xl/workbook.xml' else original.read(name))
    error = cli('submit', '--root', root, '--client', 'alice', '--task', 'week1',
                '--category', 'orders', '--period', '2026-09', '--input', path,
                '--schema', 'config/schema.example.json', code=2)
    assert json.loads(error) == {'code': 'INVALID_INPUT', 'task': 'week1'}


def test_unexpected_failure_is_redacted(tmp_path, monkeypatch, capsys):
    import service_b10.cli as module
    def broken(_):
        raise RuntimeError('sk-secret /private/customer/path')
    monkeypatch.setattr(module, '_run', broken)
    assert module.main(['init', '--root', str(tmp_path)]) == 1
    assert json.loads(capsys.readouterr().err) == {'code': 'INTERNAL_ERROR'}


@pytest.mark.parametrize('existing_draft', [False, True], ids=['first', 'amended'])
@pytest.mark.parametrize('failure', ['event', 'commit'])
def test_produce_atomic_snapshot_and_submit_rollback(tmp_path, monkeypatch, capsys,
                                                    existing_draft, failure):
    """Event/commit failure must not leave a new or overwritten workflow draft."""
    import sqlite3
    import service_b10.cli as module
    from service_b10.inputs import import_file
    from service_b10.reporting import calculate_batches
    from service_b10.inputs import PeriodImports
    from service_b10.workflow import stage_revision, record_event, freeze_metric_rules

    store = Store(tmp_path / 'live')
    store.create('alice', 'week1')
    evidence = store.put('alice', 'week1', 'proof.txt', b'synthetic reviewer evidence')
    actor = {'name': 'reviewer', 'channel': 'operator', 'evidence': evidence}
    actor_path = write_json(tmp_path / 'actor.json', actor)
    schema = json.loads(Path('config/schema.example.json').read_text())
    original_bytes = '店铺,订单号,实收\na,1,1000.00\n'.encode()
    original = import_file(store, 'alice', 'week1', 'orders', '2026-09', 'orders.csv',
                           original_bytes, schema['orders'])
    refunds = import_file(store, 'alice', 'week1', 'refunds', '2026-09', 'refunds.csv',
                          '退款号,店铺,订单号,退款金额\nr1,a,1,100.00\n'.encode(), schema['refunds'])
    if existing_draft:
        report = calculate_batches(PeriodImports(original, refunds, None))
        stage_revision(store, 'alice', 'week1',
                       {key: report[key] for key in ('orders', 'refunds', 'ads')}, report)
        for action in ('submit', 'review'):
            record_event(store, 'alice', 'week1', 1, action, actor, '')
        freeze_metric_rules(store, 'alice', 'week1', 1, {'rule': '人工核对口径'}, actor)
        record_event(store, 'alice', 'week1', 1, 'deliver', actor, '')
        record_event(store, 'alice', 'week1', 1, 'return', actor, '补充漏单')
        record_event(store, 'alice', 'week1', 1, 'amend', actor, '新版本重新生产')
        replacement = import_file(store, 'alice', 'week1', 'orders', '2026-09', 'new.csv',
                                  '店铺,订单号,实收\na,1,2000.00\n'.encode(), schema['orders'],
                                  supersedes=original.batch_id)
        old_revision = get_revision(store, 'alice', 'week1', 1)
        old_draft = get_revision(store, 'alice', 'week1', 2)
    else:
        with pytest.raises(ValueError, match='NOT_FOUND'):
            get_revision(store, 'alice', 'week1')

    def workflow_rows():
        with sqlite3.connect(store.db_path) as connection:
            return {table: connection.execute(f'SELECT * FROM {table} ORDER BY rowid').fetchall()
                    for table in ('workflow_revisions', 'workflow_events')}

    before = workflow_rows()
    fired = []
    if failure == 'event':
        with sqlite3.connect(store.db_path) as connection:
            connection.execute("""CREATE TRIGGER fail_submit BEFORE INSERT ON workflow_events
                WHEN NEW.action = 'submit' BEGIN
                    SELECT RAISE(ABORT, 'synthetic event failure');
                END""")
    else:
        # Deny actual SQLite COMMIT only after a new submit event was inserted.
        # The connection context manager must then roll back the entire operation.
        class CommitFailure(sqlite3.Connection):
            reject_commit = False
            def execute(self, sql, parameters=()):
                result = super().execute(sql, parameters)
                if 'INSERT INTO workflow_events' in sql:
                    self.reject_commit = True
                return result
        def failing_connect(self):
            connection = sqlite3.connect(self.db_path, factory=CommitFailure)
            connection.execute('PRAGMA foreign_keys = ON')
            def authorize(action, first, second, database, trigger):
                if action == sqlite3.SQLITE_TRANSACTION and first == 'COMMIT' and connection.reject_commit:
                    fired.append('COMMIT')
                    return sqlite3.SQLITE_DENY
                return sqlite3.SQLITE_OK
            connection.set_authorizer(authorize)
            return connection

    argv = list(map(str, ['produce', '--root', store.root, '--client', 'alice', '--task', 'week1',
                         '--period', '2026-09', '--actor-record', actor_path]))
    with monkeypatch.context() as patch:
        if failure == 'commit':
            patch.setattr(Store, '_connect', failing_connect)
        assert module.main(argv) == 1
    assert json.loads(capsys.readouterr().err) == {'code': 'IO_ERROR', 'task': 'week1'}
    if failure == 'commit':
        assert fired == ['COMMIT']
    assert workflow_rows() == before
    assert store.read('alice', 'week1', original.digest) == original_bytes
    if existing_draft:
        assert get_revision(store, 'alice', 'week1', 1) == old_revision
        assert get_revision(store, 'alice', 'week1', 2) == old_draft
        assert get_active_import(store, 'alice', 'week1', 'orders', '2026-09').batch_id == replacement.batch_id
    else:
        with pytest.raises(ValueError, match='NOT_FOUND'):
            get_revision(store, 'alice', 'week1')
    if failure == 'event':
        with sqlite3.connect(store.db_path) as connection:
            connection.execute('DROP TRIGGER fail_submit')
    assert module.main(argv) == 0
    snapshot = get_revision(store, 'alice', 'week1')
    assert snapshot['revision'] == (2 if existing_draft else 1)
    assert snapshot['state'] == 'submitted'
    assert snapshot['report_snapshot']['paid_less_refund'] == (190000 if existing_draft else 90000)
    assert [event['action'] for event in snapshot['events']] == ['submit']
